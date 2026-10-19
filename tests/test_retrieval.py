import json
import sqlite3
import threading
import time
from pathlib import Path
from unittest.mock import patch

import numpy as np
import pytest
from docx import Document
from langchain_core.messages import AIMessage, AIMessageChunk, HumanMessage, ToolMessage
from langgraph.checkpoint.memory import InMemorySaver
from langgraph.graph import END, START, StateGraph
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from src.backend.agent.context import AgentContext
from src.backend.agent.nodes import approval_node, retrieval_node, route_after_approval, tool_node
from src.backend.agent.state import AgentState
from src.backend.services.cancellation import CancellationToken, RunCancelledError
from src.backend.services.document_parser import parse_document
from src.backend.services.embedding import E5Encoder
from src.backend.services.retrieval import RetrievalService


@pytest.fixture(scope="module")
def encoder():
    model = Path(__file__).resolve().parents[1] / "src/backend/data/models/multilingual-e5-small"
    if not (model / "model.onnx").exists():
        pytest.skip("真实 E5 资源未下载")
    return E5Encoder(model)


@pytest.fixture
def service(tmp_path, encoder):
    item = RetrievalService(tmp_path / "index.sqlite3", encoder, debounce_seconds=0, poll_seconds=60)
    yield item
    item.close()


def test_five_formats_and_original_locations(tmp_path, service):
    root = tmp_path / "workspace"
    root.mkdir()
    doc = Document()
    doc.add_paragraph("供应商集中增加采购中断风险，建议建立第二供应商。")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "采购风险表"
    doc.save(root / "supplier.docx")
    pdf = canvas.Canvas(str(root / "backup.pdf"))
    pdf.drawString(40, 760, "Database backup and disaster recovery. Daily incremental backups protect databases.")
    pdf.save()
    book = Workbook()
    book.active.title = "预算"
    book.active.append(["广告渠道", "金额"])
    book.active.append(["搜索广告", 120000])
    book.save(root / "budget.xlsx")
    (root / "leave.txt").write_text("员工带薪年假需要经理审批。", encoding="utf-8")
    (root / "incident.md").write_text("Reset passwords and revoke sessions after suspicious sign-ins.", encoding="utf-8")
    before = {p: p.read_bytes() for p in root.iterdir()}
    result = service.search("database backup recovery", [root], top_k=5)
    assert result["index"]["ready"] == 5
    assert result["files"][0]["filename"] == "backup.pdf"
    assert {Path(hit["path"]).suffix for hit in result["files"]} == {".pdf", ".docx", ".xlsx", ".txt", ".md"}
    for hit in result["files"]:
        path = Path(hit["path"])
        assert path.read_bytes() == before[path]
        blocks = parse_document(path)
        for match in hit["matches"]:
            location = match["location"]
            block = next(b for b in blocks if all(location.get(k) == v for k, v in b.location.items()))
            assert block.text[location["char_start"]:location["char_end"]].startswith(match["text"])
    assert parse_document(root / "supplier.docx")[-1].location == {"table": 1, "row": 1}


def test_long_text_no_truncation_and_vector_normalization(tmp_path, encoder):
    path = tmp_path / "long.md"
    path.write_text("Database backups preserve information after disk failure. " * 180 + "END_MARKER", encoding="utf-8")
    chunks = encoder.split(parse_document(path))
    assert len(chunks) > 1
    assert "END_MARKER" in chunks[-1].text
    vectors = encoder.encode([c.text for c in chunks])
    assert vectors.shape == (len(chunks), 384)
    assert np.allclose(np.linalg.norm(vectors, axis=1), 1, atol=1e-5)
    assert np.isfinite(vectors).all()


def test_scope_filter_before_ranking_and_symlink_exclusion(tmp_path, service):
    a, b = tmp_path / "a", tmp_path / "b"
    a.mkdir(); b.mkdir()
    (a / "allowed.txt").write_text("Annual employee leave.")
    (b / "outside.txt").write_text("Database backups and disaster recovery.")
    (a / "link.txt").symlink_to(b / "outside.txt")
    service.synchronize([a, b])
    result = service.search("database backups", [a])
    assert [hit["filename"] for hit in result["files"]] == ["allowed.txt"]


def test_update_deletion_restart_and_unchanged_reuse(tmp_path, service, encoder):
    root = tmp_path / "workspace"; root.mkdir()
    path = root / "notes.txt"; path.write_text("Database backups protect company records.")
    original = service.search("database backups", [root])["files"][0]["sha256"]
    with patch.object(encoder, "encode", wraps=encoder.encode) as encode:
        service.synchronize([root])
        assert encode.call_count == 0
    path.write_text("Employee annual leave requires manager approval.")
    assert service.search("annual leave", [root])["index"]["pending"] == 1
    updated = service.search("annual leave", [root])
    assert updated["files"][0]["sha256"] != original
    assert "annual leave" in updated["files"][0]["matches"][0]["text"]
    service.close()
    restarted = RetrievalService(service.database, encoder, debounce_seconds=0)
    try:
        with patch.object(encoder, "encode", wraps=encoder.encode) as encode:
            restarted.synchronize([root])
            assert encode.call_count == 0
        path.unlink()
        assert restarted.search("annual leave", [root])["files"] == []
    finally:
        restarted.close()


def test_edit_during_encoding_never_publishes_old_version(tmp_path, service, encoder):
    root = tmp_path / "workspace"; root.mkdir()
    path = root / "notes.txt"; path.write_text("Original database backup policy.")
    encode = encoder.encode
    def edit_after_encoding(*args, **kwargs):
        vectors = encode(*args, **kwargs)
        path.write_text("New employee leave policy.")
        return vectors
    with patch.object(encoder, "encode", side_effect=edit_after_encoding):
        service.synchronize([root])
    with sqlite3.connect(service.database) as db:
        assert db.execute("select count(*) from chunks").fetchone()[0] == 0
        assert db.execute("select status from files").fetchone()[0] == "pending"


def test_bad_file_isolation_and_recovery(tmp_path, service):
    root = tmp_path / "workspace"; root.mkdir()
    broken = root / "broken.pdf"; broken.write_bytes(b"not a PDF")
    (root / "good.md").write_text("Database backups and recovery.")
    result = service.search("database backups", [root])
    assert result["index"]["ready"] == 1
    assert len(result["index"]["errors"]) == 1
    pdf = canvas.Canvas(str(broken)); pdf.drawString(40, 760, "Database recovery instructions."); pdf.save()
    service.synchronize([root]); service.synchronize([root])
    assert service.search("database recovery", [root])["index"]["errors"] == []


def test_model_version_change_rebuilds_vectors(tmp_path, service, encoder):
    root = tmp_path / "workspace"; root.mkdir()
    (root / "notes.txt").write_text("Database backup instructions.")
    service.synchronize([root])
    old = encoder.fingerprint
    try:
        encoder._fingerprint = old + ":new-version"
        with patch.object(encoder, "encode", wraps=encoder.encode) as encode:
            service.synchronize([root])
            assert encode.call_count == 1
    finally:
        encoder._fingerprint = old


def test_background_coalesces_saves_and_rechecks_on_start(tmp_path, encoder):
    root = tmp_path / "workspace"; root.mkdir()
    path = root / "notes.txt"; path.write_text("Original backup policy.")
    service = RetrievalService(tmp_path / "index.sqlite3", encoder, debounce_seconds=.15, poll_seconds=.02)
    service.synchronize([root])
    service.start(lambda: [str(root)])
    try:
        with patch.object(encoder, "encode", wraps=encoder.encode) as encode:
            for i in range(4):
                path.write_text(f"Revised backup policy version {i}.")
                time.sleep(.04)
            deadline = time.monotonic() + 5
            while time.monotonic() < deadline:
                with sqlite3.connect(service.database) as db:
                    texts = db.execute("select text from chunks").fetchall()
                    status = db.execute("select status from files").fetchone()[0]
                if status == "ready" and texts and "version 3" in texts[0][0]:
                    break
                time.sleep(.03)
            else:
                pytest.fail("后台索引未同步最新版本")
            assert encode.call_count == 1
    finally:
        service.close()
    path.write_text("Offline modification after application shutdown.")
    restarted = RetrievalService(service.database, encoder, debounce_seconds=.02, poll_seconds=.02)
    restarted.start(lambda: [str(root)])
    try:
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            with sqlite3.connect(service.database) as db:
                texts = db.execute("select text from chunks").fetchall()
            if texts and texts[0][0].startswith("Offline modification"):
                break
            time.sleep(.03)
        else:
            pytest.fail("启动补查未更新索引")
    finally:
        restarted.close()


def test_cancelled_search_does_not_index(tmp_path, service):
    token = CancellationToken(); token.cancel()
    with pytest.raises(RunCancelledError):
        service.search("backups", [tmp_path], cancel_check=token.raise_if_cancelled)


def test_cancel_while_waiting_for_background_index_lock(tmp_path, service):
    token = CancellationToken()
    entered = threading.Event()
    errors = []
    def run():
        entered.set()
        try:
            service.search("backups", [tmp_path], cancel_check=token.raise_if_cancelled)
        except RunCancelledError as error:
            errors.append(error)
    with service._lock:
        worker = threading.Thread(target=run)
        worker.start()
        assert entered.wait(1)
        token.cancel()
        worker.join(timeout=1)
        assert not worker.is_alive()
        assert len(errors) == 1


def test_graph_routes_real_retrieval_and_direct_read_separately(tmp_path, service):
    root = tmp_path / "workspace"; root.mkdir()
    (root / "backup.txt").write_text("Database backups protect company records.")
    builder = StateGraph(AgentState, context_schema=AgentContext)
    builder.add_node("approval", approval_node)
    builder.add_node("retrieval", retrieval_node)
    builder.add_node("tools", tool_node)
    builder.add_edge(START, "approval")
    builder.add_conditional_edges("approval", route_after_approval, {"retrieval": "retrieval", "tools": "tools", "document": "tools"})
    builder.add_edge("retrieval", END); builder.add_edge("tools", END)
    graph = builder.compile()
    context = AgentContext(user_id="test", workspace_path=str(root), retrieval_service=service)
    for name, args, expected in [
        ("search_workspace_documents", {"query": "database backups"}, "retrieval"),
        ("read_text_file", {"relative_path": "backup.txt"}, "tools"),
    ]:
        state = {"messages": [AIMessage(id="assistant", content="", tool_calls=[{"id": "call", "name": name, "args": args}])], "attachments": {}, "artifacts": {}}
        events = list(graph.stream(state, context=context, stream_mode="updates"))
        assert [next(iter(e)) for e in events] == ["approval", expected]
        payload = json.loads(events[-1][expected]["messages"][0].content)
        assert payload["ok"] is True
        if expected == "retrieval":
            assert payload["data"]["files"][0]["path"] == str((root / "backup.txt").resolve())


@pytest.mark.parametrize("request_retrieval", [True, False])
@pytest.mark.parametrize("external_source", [True, False])
def test_production_graph_turn_with_real_retrieval(tmp_path, service, request_retrieval, external_source):
    from types import SimpleNamespace
    from src.backend.agent.graph import builder

    root = tmp_path / "workspace"; root.mkdir()
    source = tmp_path / "backup.txt" if external_source else root / "backup.txt"
    source.write_text("Daily database backups protect records after disk failure.")
    graph = builder.compile(checkpointer=InMemorySaver())
    context = AgentContext(user_id="test", workspace_path=str(root), retrieval_service=service,
                           accessible_paths=(str(root), str(source)) if external_source else ())
    calls = []
    class FixedChatModel:
        def stream(self, messages, config=None):
            calls.append(messages)
            evidence = [m for m in messages if isinstance(m, ToolMessage)]
            if not request_retrieval:
                yield AIMessageChunk(content="你好", id="reply")
            elif not evidence:
                yield AIMessageChunk(content="", id="find", tool_calls=[{"id": "retrieve", "name": "search_workspace_documents", "args": {"query": "database backups"}}])
            else:
                payload = json.loads(evidence[-1].content)
                assert payload["ok"]
                yield AIMessageChunk(content="找到文件：" + payload["data"]["files"][0]["path"], id="reply")
    model = FixedChatModel()
    resolved = SimpleNamespace(model=model, model_with_tools=model, provider=context.provider, model_name=context.model_name,
                               summarizer=SimpleNamespace(before_model=lambda state, runtime: {}))
    state = {"messages": [HumanMessage(content="查找备份相关文件" if request_retrieval else "你好", id="user")], "attachments": {}, "artifacts": {},
             "prepare": {"turn_number": 1, "input_message_id": "user", "selected_attachment_ids": []}, "tool_execution_plan": None}
    with patch("src.backend.agent.nodes.model_resolver.resolve", return_value=resolved):
        events = list(graph.stream(state, {"configurable": {"thread_id": "real-retrieval-test"}}, context=context, stream_mode="updates"))
    nodes = [next(iter(e)) for e in events]
    assert ("retrieval" in nodes) == request_retrieval
    assert nodes[-1] == "finalize"
    if request_retrieval:
        assert len(calls) == 2
        assert str(source) in events[-1]["finalize"]["finalization"]["message"]["content"]
    else:
        assert not service.database.exists()


def test_legacy_imports_searchable_without_indexing_private_metadata(tmp_path):
    root = tmp_path.resolve()
    legacy = root / ".documate" / "documents"
    legacy.mkdir(parents=True)
    source = legacy / "backup.txt"
    source.write_text("数据库备份：每天增量，每周完整。", encoding="utf-8")
    (root / ".documate" / "private.txt").write_text("数据库备份：私人记录", encoding="utf-8")
    service = RetrievalService(database=root / "index.sqlite3")
    try:
        result = service.search("数据库备份", [root])
        assert result["index"]["ready"] == 1
        assert result["files"][0]["path"] == str(source)
    finally:
        service.close()


def test_selected_file_roots_do_not_include_siblings_and_track_changes(tmp_path, service):
    root = tmp_path.resolve()
    selected, sibling = root / "selected.txt", root / "sibling.txt"
    selected.write_text("Database backups run daily. Recovery copies are verified.", encoding="utf-8")
    sibling.write_text("Database backups run hourly. This file is not selected.", encoding="utf-8")
    result = service.search("database backups", [selected])
    assert [hit["path"] for hit in result["files"]] == [str(selected)]
    assert result["index"]["ready"] == 1
    assert service.index_status([selected])["ready"] == 1
    with sqlite3.connect(service.database) as db:
        assert db.execute("SELECT path FROM files").fetchall() == [(str(selected),)]
    selected.write_text("Database backups now run weekly.", encoding="utf-8")
    assert service.index_status([selected])["pending"] == 1
    service.synchronize([selected]); service.synchronize([selected])
    assert service.index_status([selected])["ready"] == 1
    selected.unlink()
    assert service.index_status([selected])["errors"][0]["status"] == "missing"
    service.synchronize([selected])
    with sqlite3.connect(service.database) as db:
        assert db.execute("SELECT COUNT(*) FROM chunks").fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM files").fetchone()[0] == 0
    selected.mkdir()
    (selected / "unexpected.txt").write_text("Do not broaden a registered file to a folder.")
    service.synchronize([selected])
    with sqlite3.connect(service.database) as db:
        assert db.execute("SELECT COUNT(*) FROM files").fetchone()[0] == 0
    import shutil
    shutil.rmtree(selected)
    selected.symlink_to(sibling)
    service.synchronize([selected])
    with sqlite3.connect(service.database) as db:
        assert db.execute("SELECT COUNT(*) FROM files").fetchone()[0] == 0
    assert sibling.is_file()


def test_index_queue_nonblocking_status_and_resume(tmp_path, service):
    source = tmp_path.resolve() / "queued.txt"
    source.write_text("Local indexing queue and document retrieval.", encoding="utf-8")
    service.request_index([source])
    deadline = time.monotonic() + 20
    while time.monotonic() < deadline and service.index_status([source])["ready"] != 1:
        time.sleep(.05)
    assert service.index_status([source])["ready"] == 1
    service.close()
    restarted = RetrievalService(service.database, service.encoder, debounce_seconds=0)
    try:
        assert restarted.index_status([source])["ready"] == 1
        assert restarted.search("local indexing", [source])["files"][0]["path"] == str(source)
    finally:
        restarted.close()


def test_index_status_stays_responsive_during_embedding(tmp_path, service):
    source = tmp_path.resolve() / "slow.txt"
    source.write_text("Document indexing progress is visible while embedding runs.")
    entered, release = threading.Event(), threading.Event()
    original = service.encoder.encode
    def paused(*args, **kwargs):
        entered.set()
        assert release.wait(10)
        return original(*args, **kwargs)
    with patch.object(service.encoder, 'encode', side_effect=paused):
        try:
            service.request_index([source])
            assert entered.wait(10)
            assert service.index_status([source])["indexing"] == 1
        finally:
            release.set()
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline and service.index_status([source])["ready"] != 1:
            time.sleep(.05)
        assert service.index_status([source])["ready"] == 1
