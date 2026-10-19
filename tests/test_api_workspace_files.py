from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pypdf import PdfWriter

from src.backend.main import create_app
from src.backend.services.app_config import AppConfigService
from tests.test_api_documents import FakeCoordinator


def test_original_tree_paging_scope_and_file_changes(tmp_path: Path):
    tmp_path = tmp_path.resolve()
    root = tmp_path / "workspace"
    root.mkdir()
    other = tmp_path / "other"
    other.mkdir()
    nested = root / "nested"
    nested.mkdir()
    (nested / "same.txt").write_text("原文件内容", encoding="utf-8")
    (other / "same.txt").write_text("另一文件", encoding="utf-8")
    for i in range(205):
        (root / f"note-{i:03}.md").write_text("测试", encoding="utf-8")
    hidden = root / ".private"
    hidden.mkdir()
    (hidden / "secret.txt").write_text("不可见", encoding="utf-8")
    (root / "link").symlink_to(other, target_is_directory=True)
    config = AppConfigService(tmp_path / "config.json")
    config.create_project("test", [str(root), str(other), str(nested)])
    with TestClient(create_app(FakeCoordinator(), config)) as client:
        assert client.get("/api/workspace-files").json()["total_files"] == 207
        page = client.get("/api/workspace-files/children", params={"path": str(root), "limit": 100}).json()
        assert page["total"] == 206
        assert len(page["entries"]) == 100
        assert page["entries"][0]["kind"] == "directory"
        rest = client.get("/api/workspace-files/children", params={"path": str(root), "offset": 100, "limit": 100}).json()
        assert len(rest["entries"]) == 100
        assert not set(e["id"] for e in page["entries"]) & set(e["id"] for e in rest["entries"])
        path = nested / "same.txt"
        original = client.get("/api/workspace-files/preview", params={"path": str(path)}).json()
        assert original["text"] == "原文件内容"
        assert original["id"] != client.get("/api/workspace-files/preview", params={"path": str(other / "same.txt")}).json()["id"]
        assert not (root / ".documate").exists()
        for denied in (tmp_path / "config.json", hidden / "secret.txt", root / "link" / "same.txt", root / "nested" / ".." / "note-000.md"):
            assert client.get("/api/workspace-files/preview", params={"path": str(denied)}).status_code == 422
        path.write_text("已更新", encoding="utf-8")
        assert client.get("/api/workspace-files/preview", params={"path": str(path)}).json()["text"] == "已更新"
        path.rename(nested / "renamed.txt")
        assert client.get("/api/workspace-files/preview", params={"path": str(path)}).status_code == 404
        listing = client.get("/api/workspace-files/children", params={"path": str(nested)}).json()
        assert listing["entries"][0]["filename"] == "renamed.txt"
        (root / "broken.docx").write_bytes(b"bad zip")
        assert client.get("/api/workspace-files/preview", params={"path": str(root / "broken.docx")}).status_code == 422
        assert client.get("/api/workspace-files").status_code == 200


def test_shared_parser_formats_and_preview_limit(tmp_path: Path):
    root = tmp_path.resolve()
    (root / "readme.md").write_text("中文说明", encoding="utf-8")
    (root / "long.txt").write_text("a" * 200_010, encoding="utf-8")
    doc = Document()
    doc.add_paragraph("合同正文")
    doc.save(root / "contract.docx")
    book = Workbook()
    book.active.append(["项目", "金额"])
    book.active.append(["测试", 123])
    book.save(root / "sheet.xlsx")
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(root / "blank.pdf")
    config = AppConfigService(root / ".config.json")
    config.switch_workspace(str(root))
    with TestClient(create_app(FakeCoordinator(), config)) as client:
        for name, expected in (("readme.md", "中文说明"), ("contract.docx", "合同正文"), ("sheet.xlsx", "123"), ("blank.pdf", "")):
            response = client.get("/api/workspace-files/preview", params={"path": str(root / name)})
            assert response.status_code == 200
            assert expected in response.json()["text"]
        long = client.get("/api/workspace-files/preview", params={"path": str(root / "long.txt")}).json()
        assert long["char_count"] == 200_010
        assert len(long["text"]) == 200_000
        assert long["truncated"]


def test_legacy_imports_visible_without_moving_and_stats_match(tmp_path: Path):
    from tests.test_api_workspace_settings import FakeCoordinator as StatsCoordinator

    root = tmp_path.resolve() / "workspace"
    root.mkdir()
    legacy = root / ".documate" / "documents"
    legacy.mkdir(parents=True)
    document = Document()
    document.add_paragraph("旧版合同仍在原位置")
    document.save(legacy / "old.docx")
    (root / "new.txt").write_text("新原文件", encoding="utf-8")
    (root / ".documate" / "private.txt").write_text("不能显示", encoding="utf-8")
    config = AppConfigService(tmp_path / "config.json")
    config.switch_workspace(str(root))
    with TestClient(create_app(StatsCoordinator(), config)) as client:
        result = client.get("/api/workspace-files").json()
        assert result["total_files"] == 2
        assert any(item["id"] == str(legacy) and "旧版导入资料" in item["filename"] for item in result["roots"])
        assert client.get("/api/stats").json()["document_count"] == result["total_files"]
        assert client.get("/api/workspace-files/preview", params={"path": str(legacy / "old.docx")}).json()["text"] == "旧版合同仍在原位置"
        assert client.get("/api/workspace-files/preview", params={"path": str(root / ".documate" / "private.txt")}).status_code == 422
    assert (legacy / "old.docx").is_file()


def test_unchanged_preview_skips_parser_and_changes_refresh(tmp_path: Path, monkeypatch):
    from src.backend.api import workspace_files

    root = tmp_path.resolve()
    source = root / "note.txt"
    source.write_text("初始内容", encoding="utf-8")
    config = AppConfigService(root / ".config.json")
    config.switch_workspace(str(root))
    actual = workspace_files.parse_document
    calls = []
    def parse(path):
        calls.append(path)
        return actual(path)
    monkeypatch.setattr(workspace_files, "parse_document", parse)
    with TestClient(create_app(FakeCoordinator(), config)) as client:
        first = client.get("/api/workspace-files/preview", params={"path": str(source)}).json()
        assert len(calls) == 1
        unchanged = client.get("/api/workspace-files/preview", params={"path": str(source), "known_version": first["version"]}).json()
        assert unchanged["unchanged"] is True
        assert len(calls) == 1
        source.write_text("修改后的新内容", encoding="utf-8")
        updated = client.get("/api/workspace-files/preview", params={"path": str(source), "known_version": first["version"]}).json()
        assert updated["text"] == "修改后的新内容"
        assert updated["version"] != first["version"]
        assert len(calls) == 2


def test_register_original_sources_scope_persistence_and_candidates(tmp_path: Path, monkeypatch):
    from src.backend.api.workspace_files import get_index
    from src.backend.api.chat import _project_context
    from src.backend.services.retrieval import RetrievalService

    root = tmp_path.resolve()
    workspace, outside = root / "workspace", root / "outside"
    workspace.mkdir(); outside.mkdir()
    selected, unselected = outside / "selected.txt", outside / "unselected.txt"
    selected.write_text("选中的资料", encoding="utf-8")
    unselected.write_text("不能自动纳入的资料", encoding="utf-8")
    (outside / ".hidden.txt").write_text("hidden")
    (outside / "unsupported.png").write_bytes(b"image")
    (outside / "linked.txt").symlink_to(selected)
    config = AppConfigService(root / "config.json")
    project = config.create_project("source test", [str(workspace)])
    index = RetrievalService(root / "index.sqlite3")
    queued = []
    monkeypatch.setattr(index, "request_index", lambda paths: queued.append(paths))
    app = create_app(FakeCoordinator(), config)
    app.dependency_overrides[get_index] = lambda: index
    body = {"project_id": project["id"], "workspace_path": str(workspace),
            "sources": [{"path": str(selected), "kind": "file"}]}
    with TestClient(app) as client:
        candidates = client.post("/api/workspace-files/candidates", json={"paths": [str(outside)], "recursive": True})
        assert candidates.status_code == 200
        assert {file["filename"] for file in candidates.json()["files"]} == {"selected.txt", "unselected.txt"}
        assert client.post("/api/workspace-files/register", json=body).status_code == 200
        assert client.post("/api/workspace-files/register", json=body).status_code == 200
        assert client.get("/api/projects").status_code == 200
        assert "document_sources" not in client.get("/api/projects").json()[0]
        assert client.post(f"/api/projects/{project['id']}/activate").status_code == 200
        listing = client.get("/api/workspace-files").json()
        assert listing["total_files"] == 1
        assert [item["id"] for item in listing["roots"]].count(str(selected)) == 1
        assert client.get("/api/workspace-files/preview", params={"path": str(selected)}).json()["text"] == "选中的资料"
        assert client.get("/api/workspace-files/preview", params={"path": str(unselected)}).status_code == 422
        assert client.get("/api/workspace-files/children", params={"path": str(outside)}).status_code == 422
        status = client.get("/api/workspace-files/index-status").json()
        assert status["pending"] == 1
        assert client.post("/api/workspace-files/refresh").json()["queued"]
        assert set(queued[-1]) == {str(workspace), str(selected)}
        assert _project_context(project["id"], str(workspace), config)[1] == (str(workspace), str(selected))
        reloaded = AppConfigService(config.path)
        assert reloaded.workspace_sources()[-1] == {"path": str(selected), "kind": "file"}
        selected.unlink()
        assert client.get("/api/workspace-files").json()["roots"][-1]["missing"]
        assert client.get("/api/workspace-files/index-status").json()["errors"][0]["status"] == "missing"
        assert client.get("/api/workspace-files/preview", params={"path": str(selected)}).status_code == 404
        assert _project_context(project["id"], str(workspace), config)[1] == (str(workspace),)
        selected.symlink_to(unselected)
        assert client.get("/api/workspace-files").json()["roots"][-1]["missing"]
        assert client.get("/api/workspace-files/preview", params={"path": str(selected)}).status_code == 422
        assert _project_context(project["id"], str(workspace), config)[1] == (str(workspace),)
        assert str(selected) not in config.index_paths()
    assert unselected.read_text(encoding="utf-8") == "不能自动纳入的资料"
    assert not (workspace / ".documate").exists()
    assert list(workspace.iterdir()) == []


def test_register_folder_future_files_and_invalid_sources(tmp_path: Path, monkeypatch):
    from src.backend.api.workspace_files import get_index
    from src.backend.services.retrieval import RetrievalService

    root = tmp_path.resolve()
    workspace, external = root / "workspace", root / "external"
    workspace.mkdir(); external.mkdir()
    nested = external / "nested"
    nested.mkdir()
    (nested / "readme.md").write_text("文件夹资料")
    config = AppConfigService(root / "config.json")
    project = config.create_project("folder test", [str(workspace)])
    index = RetrievalService(root / "index.sqlite3")
    monkeypatch.setattr(index, "request_index", lambda paths: None)
    app = create_app(FakeCoordinator(), config)
    app.dependency_overrides[get_index] = lambda: index
    body = {"project_id": project["id"], "workspace_path": str(workspace),
            "sources": [{"path": str(external), "kind": "directory"}]}
    with TestClient(app) as client:
        assert client.post("/api/workspace-files/register", json=body).status_code == 200
        (nested / "new.txt").write_text("后续新增资料")
        assert client.get("/api/workspace-files").json()["total_files"] == 2
        assert client.get("/api/workspace-files/children", params={"path": str(nested)}).json()["total"] == 2
        before = config.path.read_bytes()
        for source in ({"path": str(root / "missing.txt"), "kind": "file"},
                       {"path": str(external), "kind": "file"},
                       {"path": str(nested / "readme.md"), "kind": "directory"}):
            assert client.post("/api/workspace-files/register", json={**body, "sources": [source]}).status_code == 422
        link = root / "link"
        link.symlink_to(external, target_is_directory=True)
        assert client.post("/api/workspace-files/register", json={**body, "sources": [{"path": str(link), "kind": "directory"}]}).status_code == 422
        assert client.post("/api/workspace-files/register", json={**body, "workspace_path": str(external)}).status_code == 422
        assert config.path.read_bytes() == before
    assert not (workspace / ".documate").exists()
