import hashlib
import mimetypes
from zipfile import BadZipFile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pypdf.errors import PdfReadError


ALLOWED_DOCUMENT_SUFFIXES = {".docx", ".pdf", ".xlsx"}
MAX_UPLOAD_BYTES = 50 * 1024 * 1024
MAX_PREVIEW_CHARS = 200_000


class DocumentLibrary:
    """当前工作区的本地文档上传、列表、预览和关键词搜索。"""

    def list_documents(self, workspace_path: str) -> list[dict[str, Any]]:
        directory = self._documents_directory(workspace_path)
        if not directory.exists():
            return []
        documents = [
            self._metadata(path)
            for path in directory.iterdir()
            if path.is_file() and path.suffix.lower() in ALLOWED_DOCUMENT_SUFFIXES
        ]
        return sorted(documents, key=lambda item: item["uploaded_at"], reverse=True)

    def upload(
        self,
        workspace_path: str,
        *,
        filename: str,
        content: bytes,
    ) -> dict[str, Any]:
        safe_name = Path(filename).name
        if not safe_name or safe_name != filename or safe_name in {".", ".."}:
            raise ValueError("文件名无效")
        suffix = Path(safe_name).suffix.lower()
        if suffix not in ALLOWED_DOCUMENT_SUFFIXES:
            raise ValueError("仅支持 DOCX、PDF 和 XLSX 文件")
        if not content:
            raise ValueError("上传文件不能为空")
        if len(content) > MAX_UPLOAD_BYTES:
            raise ValueError("上传文件不能超过 50 MB")

        directory = self._documents_directory(workspace_path)
        directory.mkdir(parents=True, exist_ok=True)
        destination = directory / safe_name
        if destination.exists():
            raise FileExistsError(f"文件已存在：{safe_name}")
        temporary = directory / f".{uuid4().hex}.{safe_name}"
        try:
            temporary.write_bytes(content)
            try:
                self._extract_text(temporary)
            except ValueError as error:
                raise ValueError(f"无法读取文档：{safe_name}") from error
            temporary.replace(destination)
        finally:
            temporary.unlink(missing_ok=True)
        return self.detail(workspace_path, self._document_id(destination))

    def detail(self, workspace_path: str, document_id: str) -> dict[str, Any]:
        path = self._find(workspace_path, document_id)
        text = self._extract_text(path)
        return {**self._metadata(path, text=text), "text": text}

    def search(
        self,
        workspace_path: str,
        *,
        query: str,
        limit: int = 20,
    ) -> list[dict[str, Any]]:
        needle = query.strip().casefold()
        if not needle:
            raise ValueError("搜索内容不能为空")
        if not 1 <= limit <= 100:
            raise ValueError("limit 必须在 1 到 100 之间")

        hits: list[dict[str, Any]] = []
        for document in self.list_documents(workspace_path):
            text = self._extract_text(Path(document["saved_path"]))
            lowered = text.casefold()
            start = 0
            index = 0
            while len(hits) < limit:
                position = lowered.find(needle, start)
                if position < 0:
                    break
                left = max(0, position - 120)
                right = min(len(text), position + len(query) + 180)
                hits.append({
                    "chunk_id": f"{document['id']}:{index}",
                    "text": text[left:right],
                    "doc_id": document["id"],
                    "doc_filename": document["filename"],
                    "chunk_index": index,
                })
                index += 1
                start = position + max(1, len(needle))
            if len(hits) >= limit:
                break
        return hits

    def _documents_directory(self, workspace_path: str) -> Path:
        workspace = Path(workspace_path).expanduser().resolve()
        if not workspace.is_dir():
            raise ValueError(f"工作区不存在或不是目录：{workspace}")
        return workspace / ".documate" / "documents"

    def _find(self, workspace_path: str, document_id: str) -> Path:
        for document in self.list_documents(workspace_path):
            if document["id"] == document_id:
                return Path(document["saved_path"])
        raise FileNotFoundError(f"找不到文档：{document_id}")

    def _metadata(self, path: Path, *, text: str | None = None) -> dict[str, Any]:
        if text is None:
            text = self._extract_text(path)
        modified = datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc)
        return {
            "id": self._document_id(path),
            "filename": path.name,
            "doc_type": path.suffix.lower().lstrip("."),
            "mime_type": mimetypes.guess_type(path.name)[0] or "application/octet-stream",
            "size_bytes": path.stat().st_size,
            "char_count": len(text),
            "uploaded_at": modified.isoformat(),
            "saved_path": str(path),
            "used_ocr": False,
        }

    def _extract_text(self, path: Path) -> str:
        suffix = path.suffix.lower()
        try:
            if suffix == ".docx":
                document = Document(path)
                values = [paragraph.text for paragraph in document.paragraphs]
                for table in document.tables:
                    values.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
                text = "\n".join(value for value in values if value)
            elif suffix == ".pdf":
                text = "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
            elif suffix == ".xlsx":
                workbook = load_workbook(path, read_only=True, data_only=True)
                try:
                    lines: list[str] = []
                    for sheet in workbook.worksheets:
                        lines.append(f"[{sheet.title}]")
                        for row in sheet.iter_rows(values_only=True):
                            lines.append("\t".join("" if value is None else str(value) for value in row))
                    text = "\n".join(lines)
                finally:
                    workbook.close()
            else:
                raise ValueError(f"不支持的文档格式：{suffix}")
        except (BadZipFile, OSError, PdfReadError, ValueError) as error:
            raise ValueError(f"无法读取文档：{path.name}") from error
        return text[:MAX_PREVIEW_CHARS]

    @staticmethod
    def _document_id(path: Path) -> str:
        return hashlib.sha256(path.name.encode("utf-8")).hexdigest()[:24]
