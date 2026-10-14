"""共用的原文件解析器；不复制原文件，不截断索引正文。"""

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.table import Table
from openpyxl import load_workbook
from pypdf import PdfReader


SUPPORTED_SUFFIXES = frozenset({".pdf", ".docx", ".xlsx", ".txt", ".md"})


@dataclass(frozen=True)
class TextBlock:
    text: str
    location: dict[str, Any]


def parse_document(path: Path, content: bytes | None = None) -> list[TextBlock]:
    """解析指定版本的文件内容，保留可核对的来源位置。"""
    content = path.read_bytes() if content is None else content
    blocks: list[TextBlock] = []
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"不支持的资料格式：{suffix}")
    if suffix == ".pdf":
        for page, item in enumerate(PdfReader(BytesIO(content)).pages, 1):
            blocks.append(TextBlock(item.extract_text() or "", {"page": page}))
    elif suffix == ".docx":
        paragraph = table = 0
        for item in Document(BytesIO(content)).iter_inner_content():
            if isinstance(item, Table):
                table += 1
                for row, cells in enumerate(item.rows, 1):
                    blocks.append(TextBlock("\t".join(c.text for c in cells.cells), {"table": table, "row": row}))
            else:
                paragraph += 1
                blocks.append(TextBlock(item.text, {"paragraph": paragraph}))
    elif suffix == ".xlsx":
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            for sheet in book:
                rows = sheet.iter_rows(values_only=True)
                header = "\t".join("" if v is None else str(v) for v in next(rows, ()))
                for row, values in enumerate(rows, 2):
                    if not any(v is not None for v in values):
                        continue
                    text = f"工作表：{sheet.title}\n表头：{header}\n" + "\t".join("" if v is None else str(v) for v in values)
                    blocks.append(TextBlock(text, {"sheet": sheet.title, "row": row, "columns": len(values)}))
        finally:
            book.close()
    else:
        for line, text in enumerate(content.decode("utf-8-sig").splitlines(), 1):
            blocks.append(TextBlock(text, {"line": line}))
    return [block for block in blocks if block.text.strip()]
