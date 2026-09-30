"""DOCX、PDF 与 XLSX 专业文档工具。

工具只向模型返回有界的结构化摘要；生成或提取的文件通过 ``artifacts``
字段交给 Agent 节点登记，不把二进制内容写入 State。
"""

import csv
import os
import shutil
import subprocess
import tempfile
from copy import copy
from pathlib import Path
from typing import Any, Literal

import pdfplumber
import pypdfium2 as pdfium
from docx import Document
from docx.shared import Inches
from langchain.tools import ToolRuntime
from langchain_core.tools import tool
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils.cell import range_boundaries
from pypdf import PdfReader, PdfWriter
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

from src.backend.agent.context import AgentContext
from src.backend.agent.path_policy import resolve_runtime_path
from src.backend.agent.state import AgentState


MAX_MODEL_TEXT_CHARS = 12_000
PDF_FONT = "STSong-Light"
pdfmetrics.registerFont(UnicodeCIDFont(PDF_FONT))


def _context_path(
    runtime: ToolRuntime[AgentContext, AgentState],
    path: str,
) -> Path:
    return resolve_runtime_path(runtime, path)


def _shown_path(
    runtime: ToolRuntime[AgentContext, AgentState],
    path: Path,
) -> str:
    workspace = Path(runtime.context.workspace_path).expanduser().resolve()
    resolved = path.resolve()
    if resolved == workspace:
        return "."
    if resolved.is_relative_to(workspace):
        return resolved.relative_to(workspace).as_posix()
    return str(resolved)


def _require_file(path: Path, suffixes: set[str]) -> None:
    if not path.exists():
        raise FileNotFoundError(f"文件不存在：{path}")
    if not path.is_file():
        raise ValueError(f"路径不是文件：{path}")
    if path.suffix.lower() not in suffixes:
        expected = "、".join(sorted(suffixes))
        raise ValueError(f"文件格式不匹配，需要：{expected}")


def _prepare_output(path: Path, overwrite: bool) -> None:
    if path.exists() and not overwrite:
        raise FileExistsError(f"目标已存在：{path}")
    if path.exists() and not path.is_file():
        raise ValueError(f"目标不是文件：{path}")
    path.parent.mkdir(parents=True, exist_ok=True)


def _atomic_save(path: Path, save: Any) -> None:
    descriptor, temporary_name = tempfile.mkstemp(
        dir=path.parent,
        prefix=f".{path.stem}-",
        suffix=path.suffix,
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        save(temporary)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def _artifact(
    runtime: ToolRuntime[AgentContext, AgentState],
    path: Path,
    operation: Literal["created", "modified"] = "created",
) -> dict[str, str]:
    return {"path": _shown_path(runtime, path), "operation": operation}


def _truncate(text: str, limit: int = MAX_MODEL_TEXT_CHARS) -> tuple[str, bool]:
    if len(text) <= limit:
        return text, False
    return text[:limit], True


def _libreoffice() -> str:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if executable is None:
        raise RuntimeError("未找到 LibreOffice/soffice，无法转换文档")
    return executable


def _convert_with_libreoffice(source: Path, extension: str) -> Path:
    with tempfile.TemporaryDirectory() as directory:
        output_directory = Path(directory)
        completed = subprocess.run(
            [
                _libreoffice(),
                "--headless",
                "--convert-to",
                extension.lstrip("."),
                "--outdir",
                str(output_directory),
                str(source),
            ],
            check=False,
            capture_output=True,
            text=True,
            timeout=120,
        )
        converted = output_directory / f"{source.stem}.{extension.lstrip('.')}"
        if completed.returncode != 0 or not converted.exists():
            detail = completed.stderr.strip() or completed.stdout.strip()
            raise RuntimeError(f"LibreOffice 转换失败：{detail or 'unknown error'}")
        descriptor, temporary_name = tempfile.mkstemp(suffix=converted.suffix)
        os.close(descriptor)
        temporary = Path(temporary_name)
        shutil.copy2(converted, temporary)
        return temporary


def _docx_paragraph_text(document: Any) -> list[str]:
    return [paragraph.text for paragraph in document.paragraphs if paragraph.text]


@tool
def read_docx(
    runtime: ToolRuntime[AgentContext, AgentState],
    relative_path: str,
    max_chars: int = MAX_MODEL_TEXT_CHARS,
) -> dict[str, Any]:
    """读取 DOCX 的正文、标题、表格和页眉页脚，并返回有界结构化内容。"""
    if not 1 <= max_chars <= 50_000:
        raise ValueError("max_chars 必须在 1 到 50000 之间")
    path = _context_path(runtime, relative_path)
    _require_file(path, {".docx"})
    document = Document(path)
    paragraphs = _docx_paragraph_text(document)
    headings = [
        {"level": paragraph.style.name, "text": paragraph.text}
        for paragraph in document.paragraphs
        if paragraph.text and paragraph.style.name.startswith("Heading")
    ]
    tables = [
        [[cell.text for cell in row.cells] for row in table.rows[:100]]
        for table in document.tables[:20]
    ]
    headers: list[str] = []
    footers: list[str] = []
    for section in document.sections:
        headers.extend(_docx_paragraph_text(section.header))
        footers.extend(_docx_paragraph_text(section.footer))
    content, truncated = _truncate("\n".join(paragraphs), max_chars)
    return {
        "path": _shown_path(runtime, path),
        "paragraph_count": len(document.paragraphs),
        "heading_count": len(headings),
        "table_count": len(tables),
        "section_count": len(document.sections),
        "headings": headings,
        "tables": tables,
        "tables_truncated": len(document.tables) > 20 or any(
            len(table.rows) > 100 for table in document.tables[:20]
        ),
        "headers": headers,
        "footers": footers,
        "content": content,
        "truncated": truncated,
    }


@tool
def create_docx(
    runtime: ToolRuntime[AgentContext, AgentState],
    relative_path: str,
    paragraphs: list[str],
    title: str | None = None,
    tables: list[list[list[str]]] | None = None,
    image_paths: list[str] | None = None,
    overwrite: bool = False,
) -> dict[str, Any]:
    """创建 DOCX，可写入标题、段落、表格和图片。"""
    path = _context_path(runtime, relative_path)
    if path.suffix.lower() != ".docx":
        raise ValueError("目标必须是 .docx 文件")
    existed = path.exists()
    _prepare_output(path, overwrite)
    document = Document()
    if title:
        document.add_heading(title, level=0)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    for rows in tables or []:
        width = max((len(row) for row in rows), default=0)
        if not rows or width == 0:
            continue
        table = document.add_table(rows=len(rows), cols=width)
        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = str(value)
    for image_path in image_paths or []:
        image = _context_path(runtime, image_path)
        _require_file(image, {".png", ".jpg", ".jpeg"})
        document.add_picture(str(image), width=Inches(6))
    _atomic_save(path, document.save)
    return {
        "path": _shown_path(runtime, path),
        "paragraph_count": len(paragraphs),
        "table_count": len(tables or []),
        "image_count": len(image_paths or []),
        "overwritten": existed,
        "artifacts": [_artifact(runtime, path, "modified" if existed else "created")],
    }


def _replace_paragraph_text(paragraph: Any, old: str, new: str) -> int:
    full_text = "".join(run.text for run in paragraph.runs)
    count = full_text.count(old)
    if count == 0:
        return 0
    replaced = full_text.replace(old, new)
    if paragraph.runs:
        paragraph.runs[0].text = replaced
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(replaced)
    return count


@tool
def edit_docx(
    runtime: ToolRuntime[AgentContext, AgentState],
    relative_path: str,
    output_path: str,
    replacements: dict[str, str] | None = None,
    append_paragraphs: list[str] | None = None,
    overwrite: bool = False,
) -> dict[str, Any]:
    """修改 DOCX 中的正文或表格文字，并可追加段落后另存。"""
    source = _context_path(runtime, relative_path)
    destination = _context_path(runtime, output_path)
    _require_file(source, {".docx"})
    if destination.suffix.lower() != ".docx":
        raise ValueError("输出必须是 .docx 文件")
    existed = destination.exists()
    _prepare_output(destination, overwrite)
    document = Document(source)
    replacement_count = 0
    for old, new in (replacements or {}).items():
        if not old:
            raise ValueError("替换的原文本不能为空")
        for paragraph in document.paragraphs:
            replacement_count += _replace_paragraph_text(paragraph, old, new)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        replacement_count += _replace_paragraph_text(paragraph, old, new)
    for paragraph in append_paragraphs or []:
        document.add_paragraph(paragraph)
    _atomic_save(destination, document.save)
    return {
        "source": _shown_path(runtime, source),
        "path": _shown_path(runtime, destination),
        "replacement_count": replacement_count,
        "appended_paragraph_count": len(append_paragraphs or []),
        "artifacts": [
            _artifact(runtime, destination, "modified" if existed else "created")
        ],
    }


@tool
def convert_docx(
    runtime: ToolRuntime[AgentContext, AgentState],
    relative_path: str,
    output_path: str,
    overwrite: bool = False,
) -> dict[str, Any]:
    """使用 LibreOffice 将 DOC/DOCX 转换为 DOCX 或 PDF。"""
    source = _context_path(runtime, relative_path)
    destination = _context_path(runtime, output_path)
    _require_file(source, {".doc", ".docx"})
    if destination.suffix.lower() not in {".docx", ".pdf"}:
        raise ValueError("输出格式只能是 .docx 或 .pdf")
    if source.suffix.lower() == destination.suffix.lower():
        raise ValueError("源格式和目标格式不能相同")
    existed = destination.exists()
    _prepare_output(destination, overwrite)
    converted = _convert_with_libreoffice(source, destination.suffix)
    try:
        os.replace(converted, destination)
    finally:
        converted.unlink(missing_ok=True)
    return {
        "source": _shown_path(runtime, source),
        "path": _shown_path(runtime, destination),
        "artifacts": [
            _artifact(runtime, destination, "modified" if existed else "created")
        ],
    }


@tool
def render_docx(
    runtime: ToolRuntime[AgentContext, AgentState],
    relative_path: str,
    output_directory: str,
    dpi: int = 120,
    overwrite: bool = False,
) -> dict[str, Any]:
    """将 DOCX 渲染为 PDF 和逐页 PNG 预览。"""
    if not 72 <= dpi <= 300:
        raise ValueError("dpi 必须在 72 到 300 之间")
    source = _context_path(runtime, relative_path)
    directory = _context_path(runtime, output_directory)
    _require_file(source, {".docx"})
    if directory.exists() and any(directory.iterdir()) and not overwrite:
        raise FileExistsError(f"输出目录非空：{directory}")
    directory.mkdir(parents=True, exist_ok=True)
    converted = _convert_with_libreoffice(source, ".pdf")
    pdf_path = directory / f"{source.stem}.pdf"
    os.replace(converted, pdf_path)
    pdf = pdfium.PdfDocument(pdf_path)
    previews: list[Path] = []
    scale = dpi / 72
    for index in range(len(pdf)):
        preview = directory / f"page-{index + 1:04d}.png"
        pdf[index].render(scale=scale).to_pil().save(preview)
        previews.append(preview)
    outputs = [pdf_path, *previews]
    return {
        "source": _shown_path(runtime, source),
        "page_count": len(previews),
        "preview_paths": [_shown_path(runtime, item) for item in previews],
        "artifacts": [_artifact(runtime, item) for item in outputs],
    }


@tool
def validate_docx(
    runtime: ToolRuntime[AgentContext, AgentState],
    relative_path: str,
    check_rendering: bool = True,
) -> dict[str, Any]:
    """验证 DOCX 能否解析，并可通过临时 PDF 转换检查渲染。"""
    path = _context_path(runtime, relative_path)
    _require_file(path, {".docx"})
    document = Document(path)
    rendered = False
    page_count: int | None = None
    if check_rendering:
        converted = _convert_with_libreoffice(path, ".pdf")
        try:
            page_count = len(PdfReader(converted).pages)
            rendered = True
        finally:
            converted.unlink(missing_ok=True)
    return {
        "path": _shown_path(runtime, path),
        "valid": True,
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "rendered": rendered,
        "page_count": page_count,
    }


@tool
def read_pdf(
    runtime: ToolRuntime[AgentContext, AgentState],
    relative_path: str,
    start_page: int = 1,
    page_count: int = 20,
    max_chars: int = MAX_MODEL_TEXT_CHARS,
) -> dict[str, Any]:
    """读取 PDF 指定页范围的文本、页码、元数据和基础结构。"""
    path = _context_path(runtime, relative_path)
    _require_file(path, {".pdf"})
    if start_page < 1 or not 1 <= page_count <= 200:
        raise ValueError("页码参数无效")
    reader = PdfReader(path)
    selected = reader.pages[start_page - 1 : start_page - 1 + page_count]
    parts = [f"[第 {start_page + index} 页]\n{page.extract_text() or ''}" for index, page in enumerate(selected)]
    content, truncated = _truncate("\n\n".join(parts), max_chars)
    metadata = reader.metadata or {}
    return {
        "path": _shown_path(runtime, path),
        "total_pages": len(reader.pages),
        "start_page": start_page,
        "returned_pages": len(selected),
        "metadata": {str(key): str(value) for key, value in metadata.items()},
        "content": content,
        "truncated": truncated or start_page - 1 + len(selected) < len(reader.pages),
    }


@tool
def extract_pdf_tables(
    runtime: ToolRuntime[AgentContext, AgentState],
    relative_path: str,
    start_page: int = 1,
    page_count: int = 20,
    max_tables: int = 20,
) -> dict[str, Any]:
    """提取 PDF 指定页范围内可识别的表格。"""
    path = _context_path(runtime, relative_path)
    _require_file(path, {".pdf"})
    if start_page < 1 or not 1 <= page_count <= 100 or not 1 <= max_tables <= 100:
        raise ValueError("页码参数无效")
    tables: list[dict[str, Any]] = []
    truncated = False
    with pdfplumber.open(path) as pdf:
        pages = pdf.pages[start_page - 1 : start_page - 1 + page_count]
        for offset, page in enumerate(pages):
            for table_index, rows in enumerate(page.extract_tables(), start=1):
                if len(tables) >= max_tables:
                    truncated = True
                    break
                tables.append(
                    {
                        "page": start_page + offset,
                        "table_index": table_index,
                        "rows": rows[:200],
                        "rows_truncated": len(rows) > 200,
                    }
                )
            if truncated:
                break
    return {
        "path": _shown_path(runtime, path),
        "table_count": len(tables),
        "tables": tables,
        "truncated": truncated,
    }


@tool
def extract_pdf_images(
    runtime: ToolRuntime[AgentContext, AgentState],
    relative_path: str,
    output_directory: str,
    overwrite: bool = False,
) -> dict[str, Any]:
    """提取 PDF 中可访问的原始图片，并登记为文件 Artifact。"""
    source = _context_path(runtime, relative_path)
    directory = _context_path(runtime, output_directory)
    _require_file(source, {".pdf"})
    if directory.exists() and any(directory.iterdir()) and not overwrite:
        raise FileExistsError(f"输出目录非空：{directory}")
    directory.mkdir(parents=True, exist_ok=True)
    reader = PdfReader(source)
    outputs: list[Path] = []
    for page_index, page in enumerate(reader.pages, start=1):
        for image_index, image in enumerate(page.images, start=1):
            suffix = Path(image.name).suffix or ".bin"
            output = directory / f"page-{page_index:04d}-image-{image_index:03d}{suffix}"
            output.write_bytes(image.data)
            outputs.append(output)
    return {
        "source": _shown_path(runtime, source),
        "image_count": len(outputs),
        "paths": [_shown_path(runtime, item) for item in outputs],
        "artifacts": [_artifact(runtime, item) for item in outputs],
    }


def _draw_wrapped_line(pdf: Any, text: str, x: float, y: float, width: float) -> float:
    words = text.split() or [""]
    line = ""
    for word in words:
        candidate = f"{line} {word}".strip()
        if line and pdfmetrics.stringWidth(candidate, PDF_FONT, 11) > width:
            pdf.drawString(x, y, line)
            y -= 16
            line = word
        else:
            line = candidate
    pdf.drawString(x, y, line)
    return y - 16


@tool
def create_pdf(
    runtime: ToolRuntime[AgentContext, AgentState],
    relative_path: str,
    paragraphs: list[str],
    title: str | None = None,
    overwrite: bool = False,
) -> dict[str, Any]:
    """从标题和段落创建基础 PDF。"""
    path = _context_path(runtime, relative_path)
    if path.suffix.lower() != ".pdf":
        raise ValueError("目标必须是 .pdf 文件")
    existed = path.exists()
    _prepare_output(path, overwrite)

    def save(output: Path) -> None:
        pdf = canvas.Canvas(str(output), pagesize=A4)
        width, height = A4
        y = height - 60
        if title:
            pdf.setFont(PDF_FONT, 18)
            pdf.drawString(50, y, title)
            y -= 30
        pdf.setFont(PDF_FONT, 11)
        for paragraph in paragraphs:
            for line in paragraph.splitlines() or [""]:
                if y < 60:
                    pdf.showPage()
                    pdf.setFont(PDF_FONT, 11)
                    y = height - 60
                y = _draw_wrapped_line(pdf, line, 50, y, width - 100)
            y -= 8
        pdf.save()

    _atomic_save(path, save)
    return {
        "path": _shown_path(runtime, path),
        "paragraph_count": len(paragraphs),
        "overwritten": existed,
        "artifacts": [_artifact(runtime, path, "modified" if existed else "created")],
    }


@tool
def edit_pdf_pages(
    runtime: ToolRuntime[AgentContext, AgentState],
    source_paths: list[str],
    output_path: str,
    page_order: list[int] | None = None,
    rotation: Literal[0, 90, 180, 270] = 0,
    overwrite: bool = False,
) -> dict[str, Any]:
    """合并 PDF，并可按一基页码提取、删除或重排页面及统一旋转。"""
    if not source_paths:
        raise ValueError("至少需要一个源 PDF")
    sources = [_context_path(runtime, item) for item in source_paths]
    for source in sources:
        _require_file(source, {".pdf"})
    destination = _context_path(runtime, output_path)
    if destination.suffix.lower() != ".pdf":
        raise ValueError("输出必须是 .pdf 文件")
    existed = destination.exists()
    _prepare_output(destination, overwrite)
    pages = [page for source in sources for page in PdfReader(source).pages]
    indexes = [number - 1 for number in page_order] if page_order else list(range(len(pages)))
    if any(index < 0 or index >= len(pages) for index in indexes):
        raise ValueError("页码超出 PDF 范围")
    writer = PdfWriter()
    for index in indexes:
        page = pages[index]
        if rotation:
            page.rotate(rotation)
        writer.add_page(page)
    _atomic_save(destination, lambda output: writer.write(str(output)))
    return {
        "sources": [_shown_path(runtime, item) for item in sources],
        "path": _shown_path(runtime, destination),
        "page_count": len(indexes),
        "artifacts": [
            _artifact(runtime, destination, "modified" if existed else "created")
        ],
    }


def _json_cell(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return value.isoformat() if hasattr(value, "isoformat") else str(value)


@tool
def read_xlsx(
    runtime: ToolRuntime[AgentContext, AgentState],
    relative_path: str,
    sheet_name: str | None = None,
    cell_range: str | None = None,
    max_rows: int = 200,
    max_columns: int = 50,
) -> dict[str, Any]:
    """读取 XLSX 的 Sheet、单元格值、公式和基础结构。"""
    path = _context_path(runtime, relative_path)
    _require_file(path, {".xlsx", ".xlsm"})
    if not 1 <= max_rows <= 2_000 or not 1 <= max_columns <= 200:
        raise ValueError("读取范围超出限制")
    workbook = load_workbook(path, data_only=False, read_only=True)
    if sheet_name and sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet 不存在：{sheet_name}")
    names = [sheet_name] if sheet_name else workbook.sheetnames[:10]
    sheets: dict[str, Any] = {}
    truncated = False
    for name in names:
        sheet = workbook[name]
        if cell_range:
            min_column, min_row, max_column, max_row = range_boundaries(cell_range)
        else:
            min_row = min_column = 1
            max_row = min(sheet.max_row, max_rows)
            max_column = min(sheet.max_column, max_columns)
            truncated = truncated or sheet.max_row > max_rows or sheet.max_column > max_columns
        rows = [
            [_json_cell(cell.value) for cell in row]
            for row in sheet.iter_rows(
                min_row=min_row,
                max_row=max_row,
                min_col=min_column,
                max_col=max_column,
            )
        ]
        sheets[name] = {
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "range": cell_range or f"A1:{sheet.cell(max_row, max_column).coordinate}",
            "rows": rows,
        }
    workbook.close()
    return {
        "path": _shown_path(runtime, path),
        "sheet_names": names,
        "all_sheet_names": workbook.sheetnames,
        "sheets_truncated": sheet_name is None and len(workbook.sheetnames) > len(names),
        "sheets": sheets,
        "truncated": truncated,
    }


@tool
def create_xlsx(
    runtime: ToolRuntime[AgentContext, AgentState],
    relative_path: str,
    sheets: dict[str, list[list[Any]]],
    overwrite: bool = False,
) -> dict[str, Any]:
    """创建 XLSX；单元格字符串以等号开头时作为公式写入。"""
    if not sheets:
        raise ValueError("至少需要一个 Sheet")
    path = _context_path(runtime, relative_path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("目标必须是 .xlsx 文件")
    existed = path.exists()
    _prepare_output(path, overwrite)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        if not name or len(name) > 31:
            raise ValueError(f"无效的 Sheet 名称：{name}")
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    _atomic_save(path, workbook.save)
    return {
        "path": _shown_path(runtime, path),
        "sheet_names": list(sheets),
        "overwritten": existed,
        "artifacts": [_artifact(runtime, path, "modified" if existed else "created")],
    }


def _apply_cell_style(cell: Any, style: dict[str, Any]) -> None:
    if "number_format" in style:
        cell.number_format = str(style["number_format"])
    if "bold" in style or "font_color" in style:
        font = copy(cell.font)
        if "bold" in style:
            font.bold = bool(style["bold"])
        if "font_color" in style:
            font.color = str(style["font_color"])
        cell.font = font
    if "fill_color" in style:
        cell.fill = PatternFill("solid", fgColor=str(style["fill_color"]))
    if "horizontal" in style:
        cell.alignment = Alignment(horizontal=str(style["horizontal"]))


@tool
def edit_xlsx(
    runtime: ToolRuntime[AgentContext, AgentState],
    relative_path: str,
    output_path: str,
    cell_updates: dict[str, dict[str, Any]] | None = None,
    add_sheets: list[str] | None = None,
    remove_sheets: list[str] | None = None,
    rename_sheets: dict[str, str] | None = None,
    cell_styles: dict[str, dict[str, dict[str, Any]]] | None = None,
    overwrite: bool = False,
) -> dict[str, Any]:
    """修改 XLSX 的单元格、Sheet 和基础样式后另存。"""
    source = _context_path(runtime, relative_path)
    destination = _context_path(runtime, output_path)
    _require_file(source, {".xlsx", ".xlsm"})
    if destination.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("输出必须是 .xlsx 或 .xlsm 文件")
    existed = destination.exists()
    _prepare_output(destination, overwrite)
    workbook = load_workbook(source, keep_vba=source.suffix.lower() == ".xlsm")
    for name in add_sheets or []:
        if name in workbook.sheetnames:
            raise ValueError(f"Sheet 已存在：{name}")
        workbook.create_sheet(name)
    for old, new in (rename_sheets or {}).items():
        if old not in workbook.sheetnames:
            raise ValueError(f"Sheet 不存在：{old}")
        workbook[old].title = new
    for name in remove_sheets or []:
        if name not in workbook.sheetnames:
            raise ValueError(f"Sheet 不存在：{name}")
        if len(workbook.sheetnames) == 1:
            raise ValueError("不能删除工作簿的最后一个 Sheet")
        workbook.remove(workbook[name])
    update_count = 0
    for sheet_name, updates in (cell_updates or {}).items():
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet 不存在：{sheet_name}")
        for coordinate, value in updates.items():
            workbook[sheet_name][coordinate] = value
            update_count += 1
    for sheet_name, styles in (cell_styles or {}).items():
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet 不存在：{sheet_name}")
        for coordinate, style in styles.items():
            _apply_cell_style(workbook[sheet_name][coordinate], style)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    _atomic_save(destination, workbook.save)
    return {
        "source": _shown_path(runtime, source),
        "path": _shown_path(runtime, destination),
        "updated_cells": update_count,
        "sheet_names": workbook.sheetnames,
        "artifacts": [
            _artifact(runtime, destination, "modified" if existed else "created")
        ],
    }


def _xlsx_to_csv(source: Path, destination: Path) -> None:
    workbook = load_workbook(source, data_only=True, read_only=True)
    sheet = workbook.active
    with destination.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerows(sheet.iter_rows(values_only=True))
    workbook.close()


def _csv_to_xlsx(source: Path, destination: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    with source.open(encoding="utf-8-sig", newline="") as stream:
        for row in csv.reader(stream):
            sheet.append(row)
    workbook.save(destination)


@tool
def convert_xlsx(
    runtime: ToolRuntime[AgentContext, AgentState],
    relative_path: str,
    output_path: str,
    overwrite: bool = False,
) -> dict[str, Any]:
    """在 XLSX、CSV 和 PDF 之间转换；PDF 转换使用 LibreOffice。"""
    source = _context_path(runtime, relative_path)
    destination = _context_path(runtime, output_path)
    _require_file(source, {".xlsx", ".csv"})
    if destination.suffix.lower() not in {".xlsx", ".csv", ".pdf"}:
        raise ValueError("输出格式只能是 .xlsx、.csv 或 .pdf")
    if source.suffix.lower() == destination.suffix.lower():
        raise ValueError("源格式和目标格式不能相同")
    if destination.suffix.lower() == ".pdf" and source.suffix.lower() != ".xlsx":
        raise ValueError("只支持 XLSX 转 PDF")
    existed = destination.exists()
    _prepare_output(destination, overwrite)
    if source.suffix.lower() == ".xlsx" and destination.suffix.lower() == ".csv":
        _atomic_save(destination, lambda output: _xlsx_to_csv(source, output))
    elif source.suffix.lower() == ".csv" and destination.suffix.lower() == ".xlsx":
        _atomic_save(destination, lambda output: _csv_to_xlsx(source, output))
    else:
        converted = _convert_with_libreoffice(source, ".pdf")
        try:
            os.replace(converted, destination)
        finally:
            converted.unlink(missing_ok=True)
    return {
        "source": _shown_path(runtime, source),
        "path": _shown_path(runtime, destination),
        "artifacts": [
            _artifact(runtime, destination, "modified" if existed else "created")
        ],
    }


@tool
def validate_xlsx(
    runtime: ToolRuntime[AgentContext, AgentState],
    relative_path: str,
) -> dict[str, Any]:
    """验证 XLSX 能否打开，并统计 Sheet、公式和已缓存错误值。"""
    path = _context_path(runtime, relative_path)
    _require_file(path, {".xlsx", ".xlsm"})
    workbook = load_workbook(path, data_only=False, read_only=True)
    cached = load_workbook(path, data_only=True, read_only=True)
    formula_count = 0
    error_cells: list[str] = []
    for sheet in workbook.worksheets:
        cached_sheet = cached[sheet.title]
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_count += 1
                    cached_value = cached_sheet[cell.coordinate].value
                    if isinstance(cached_value, str) and cached_value.startswith("#"):
                        error_cells.append(f"{sheet.title}!{cell.coordinate}")
    result = {
        "path": _shown_path(runtime, path),
        "valid": True,
        "sheet_names": workbook.sheetnames,
        "formula_count": formula_count,
        "formula_error_cells": error_cells[:100],
        "formula_errors_truncated": len(error_cells) > 100,
    }
    workbook.close()
    cached.close()
    return result


DOCUMENT_READ_ONLY_TOOLS = [
    read_docx,
    validate_docx,
    read_pdf,
    extract_pdf_tables,
    read_xlsx,
    validate_xlsx,
]

DOCUMENT_MUTATING_TOOLS = [
    create_docx,
    edit_docx,
    convert_docx,
    render_docx,
    extract_pdf_images,
    create_pdf,
    edit_pdf_pages,
    create_xlsx,
    edit_xlsx,
    convert_xlsx,
]

DOCUMENT_TOOLS = DOCUMENT_READ_ONLY_TOOLS + DOCUMENT_MUTATING_TOOLS
DOCUMENT_TOOL_NAMES = frozenset(tool.name for tool in DOCUMENT_TOOLS)
